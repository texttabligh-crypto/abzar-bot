import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func

from app.models import User, Product, Order, Payment, Category, Brand, CartItem
from app.bot.routers import payments, admin
from app.services.export import build_price_list, build_price_list_pdf
from app.services.orders import get_or_create_cart, create_order_from_cart, CheckoutError
from tests.fakes import FakeUser, FakeMessage, FakeCallback, make_state


async def seed(session):
    user = User(telegram_id=9001, first_name='Customer', phone='0912',
                customer_type='retail', is_active=True, is_approved=True)
    product = Product(internal_sku='REG-1', name='Product', base_price=100,
                      wholesale_price=70, stock_quantity=5)
    session.add_all([user, product])
    await session.commit()
    return user, product


@pytest.mark.asyncio
async def test_receipt_callback_uses_customer_not_bot_and_blocks_duplicate(db, monkeypatch):
    async with db() as s:
        user, _ = await seed(s)
        order = Order(order_no='REG-ORDER', user_id=user.id, status='WAITING_PAYMENT', total_amount=100)
        s.add(order)
        await s.commit()
        oid = order.id

    class Redis:
        deleted = False
        async def get(self, key): return 'receipt-file'
        async def delete(self, key): self.deleted = True

    import app.core.redis
    cache = Redis()
    monkeypatch.setattr(app.core.redis, 'redis_client', cache)
    message = FakeMessage(user=FakeUser(id=123, is_bot=True))
    cb = FakeCallback(f'payorder:{oid}', user=FakeUser(id=9001), message=message)
    await payments.choose_receipt_order(cb)
    assert cache.deleted
    assert 'فیش ثبت شد.' in cb.alerts
    assert not await payments.register_receipt(message, oid, 'other', customer_tg_id=9001)
    async with db() as s:
        assert await s.scalar(select(func.count(Payment.id))) == 1
        assert (await s.get(Order, oid)).status == 'PAYMENT_REVIEW'


@pytest.mark.asyncio
async def test_receipt_wrong_owner_not_successful_and_cache_preserved(db, monkeypatch):
    async with db() as s:
        user, _ = await seed(s)
        s.add(User(telegram_id=9002, first_name='Other', phone='0912', is_active=True))
        order = Order(order_no='OTHER-ORDER', user_id=user.id, status='WAITING_PAYMENT', total_amount=100)
        s.add(order)
        await s.commit()
        oid = order.id
    class Redis:
        async def get(self, key): return 'file'
        async def delete(self, key): raise AssertionError('Must preserve receipt on failure')
    import app.core.redis
    monkeypatch.setattr(app.core.redis, 'redis_client', Redis())
    cb = FakeCallback(f'payorder:{oid}', user=FakeUser(id=9002))
    await payments.choose_receipt_order(cb)
    assert 'فیش ثبت شد.' not in cb.alerts
    async with db() as s:
        assert await s.scalar(select(func.count(Payment.id))) == 0


@pytest.mark.asyncio
async def test_xlsx_hides_wholesale_for_retail(db):
    async with db() as s:
        await seed(s)
        retail = load_workbook(await build_price_list(s, include_wholesale=False)).active
        wholesale = load_workbook(await build_price_list(s, include_wholesale=True)).active
        assert retail.max_column == 7
        assert all('عمده' not in str(c.value) for c in retail[1])
        assert 70 not in [c.value for c in retail[2]]
        assert wholesale.max_column == 8
        assert wholesale.cell(2, 5).value == 70


@pytest.mark.asyncio
async def test_pdf_removes_wholesale_column(db, monkeypatch):
    import app.services.export as export
    real_table = export.Table
    captured = []
    def capture(data, **kwargs):
        captured.append(data)
        return real_table(data, **kwargs)
    monkeypatch.setattr(export, 'Table', capture)
    async with db() as s:
        await seed(s)
        assert (await build_price_list_pdf(s, include_wholesale=False)).read(4) == b'%PDF'
        assert all(len(row) == 6 for row in captured[-1])
        assert (await build_price_list_pdf(s, include_wholesale=True)).read(4) == b'%PDF'
        assert all(len(row) == 7 for row in captured[-1])


@pytest.mark.asyncio
@pytest.mark.parametrize('change', ['inactive', 'unavailable', 'out_of_stock', 'insufficient'])
async def test_checkout_revalidates_cart_and_preserves_on_failure(db, change):
    async with db() as s:
        user, product = await seed(s)
        cart = await get_or_create_cart(s, user.id)
        s.add(CartItem(cart_id=cart.id, product_id=product.id, quantity=3))
        if change == 'inactive': product.is_active = False
        elif change == 'insufficient': product.stock_quantity = 2
        else: product.stock_status = change
        await s.commit()
        with pytest.raises(CheckoutError):
            await create_order_from_cart(s, user.id)
        await s.rollback()
        assert await s.scalar(select(func.count(Order.id))) == 0
        assert await s.scalar(select(func.count(CartItem.id))) == 1


@pytest.mark.asyncio
async def test_category_price_scope_includes_descendants(db):
    async with db() as s:
        user, product = await seed(s)
        user.is_owner = True
        parent = Category(name='Parent')
        s.add(parent)
        await s.flush()
        child = Category(name='Child', parent_id=parent.id)
        s.add(child)
        await s.flush()
        product.category_id = child.id
        await s.commit()
        cid, pid = parent.id, product.id
    state = make_state(9001)
    await state.clear()
    await state.update_data(price_change_quick=True, price_change_mode='percent', price_change_field='all')
    await admin.price_change_category_pick(FakeCallback(f'pricechg:category:{cid}', user=FakeUser(id=9001)), state)
    assert pid in (await state.get_data())['selected_product_ids']


@pytest.mark.asyncio
async def test_edit_existing_product_category_and_brand(db):
    async with db() as s:
        user, product = await seed(s)
        user.is_owner = True
        cat, brand = Category(name='New category'), Brand(name='New brand')
        s.add_all([cat, brand])
        await s.commit()
        pid, cid, bid = product.id, cat.id, brand.id
    state = make_state(9001)
    await state.clear()
    await state.update_data(product_id=pid)
    for kind, oid in [('category', cid), ('brand', bid)]:
        cb = FakeCallback(f'editmeta:set:{kind}:{oid}', user=FakeUser(id=9001))
        await admin.edit_metadata_save(cb, state)
    async with db() as s:
        product = await s.get(Product, pid)
        assert product.category_id == cid and product.brand_id == bid


@pytest.mark.asyncio
@pytest.mark.parametrize('kind,approved,expected', [('retail', True, False), ('wholesale', True, True)])
async def test_customer_exports_pass_segment_policy(db, monkeypatch, kind, approved, expected):
    from io import BytesIO
    from app.bot.routers import customer
    async with db() as s:
        user, _ = await seed(s)
        user.customer_type = kind
        user.is_approved = approved
        await s.commit()
    calls = []
    async def export(session, include_wholesale=True):
        calls.append(include_wholesale)
        return BytesIO(b'test')
    monkeypatch.setattr(customer, 'build_price_list', export)
    monkeypatch.setattr(customer, 'build_price_list_pdf', export)
    cb = FakeCallback('pricefmt:xlsx', user=FakeUser(id=9001))
    await customer.price_xlsx(cb)
    await customer.price_pdf(cb)
    assert calls == [expected, expected]


@pytest.mark.asyncio
async def test_pending_wholesale_cannot_export(db, monkeypatch):
    from app.bot.routers import customer
    async with db() as s:
        user, _ = await seed(s)
        user.customer_type, user.is_approved = 'wholesale', False
        await s.commit()
    async def forbidden(*args, **kwargs):
        raise AssertionError('Export must not run for pending wholesale customer')
    monkeypatch.setattr(customer, 'build_price_list', forbidden)
    monkeypatch.setattr(customer, 'build_price_list_pdf', forbidden)
    cb = FakeCallback('pricefmt:xlsx', user=FakeUser(id=9001))
    await customer.price_xlsx(cb)
    await customer.price_pdf(cb)
    assert 'document' not in cb.message.answers
