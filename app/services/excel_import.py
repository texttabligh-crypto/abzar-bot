from typing import Any
from openpyxl import load_workbook
from rapidfuzz import fuzz

from app.utils import norm_text

ALIASES = {
    "name": ["نام", "نام کالا", "شرح", "شرح کالا", "محصول", "product", "item", "description", "product name"],
    "price": ["قیمت", "مبلغ", "فی", "قیمت فروش", "price", "sale price", "unit price", "amount"],
    "sku": ["کد", "کد کالا", "sku", "part no", "part number", "item code", "code"],
    "brand": ["برند", "brand", "manufacturer"],
}

def norm(v: Any) -> str:
    return norm_text(str(v or "").replace("_", " "))

def detect_header_row(ws, scan_rows=10):
    best_row, best_score = 1, -1
    for r in range(1, min(ws.max_row, scan_rows) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        score = 0
        for v in vals:
            for aliases in ALIASES.values():
                score += max([fuzz.ratio(v, norm(a)) for a in aliases] + [0])
        if score > best_score:
            best_row, best_score = r, score
    return best_row

def detect_mapping(headers):
    mapping = {}
    for idx, h in enumerate(headers, start=1):
        hn = norm(h)
        best_field, best = None, 0
        for field, aliases in ALIASES.items():
            score = max([fuzz.ratio(hn, norm(a)) for a in aliases] + [0])
            if score > best:
                best_field, best = field, score
        if best_field and best >= 65 and best_field not in mapping:
            mapping[best_field] = idx
    return mapping

def parse_workbook(path: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    header_row = detect_header_row(ws)
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    mapping = detect_mapping(headers)
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        row = {field: ws.cell(r, c).value for field, c in mapping.items()}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return {
        "header_row": header_row,
        "headers": headers,
        "mapping": mapping,
        "rows": rows,
    }
